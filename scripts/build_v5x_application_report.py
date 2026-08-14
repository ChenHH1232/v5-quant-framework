from __future__ import annotations

import csv
import hashlib
import json
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "v5x_personal_project_report" / "current"
FIG_V = ROOT / "v5v_application_final_report_readiness" / "current" / "figures"
FIG_W = ROOT / "v5w_final_report_preparation" / "current" / "figures"
RUNTIME_PY = Path(r"C:\Users\Administrator\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe")
DOC_RENDER = Path(r"C:\Users\Administrator\.codex\plugins\cache\openai-primary-runtime\documents\26.805.11740\skills\documents\render_docx.py")

BASELINE = "v57f_startup_preload_repaired_baseline"
CANDIDATE = "internal_subsleeve_mom12_70_30"
FORMAL_START, FORMAL_END, OBS = "2021-05-06", "2026-05-29", 1228


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def xml_tag(name: str, **attrs: str) -> OxmlElement:
    element = OxmlElement(name)
    for key, value in attrs.items():
        element.set(qn(key), value)
    return element


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def borders(cell, color: str = "D7E4EF") -> None:
    props = cell._tc.get_or_add_tcPr()
    cell_borders = props.first_child_found_in("w:tcBorders")
    if cell_borders is None:
        cell_borders = OxmlElement("w:tcBorders")
        props.append(cell_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = cell_borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            cell_borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "5")
        tag.set(qn("w:color"), color)


def set_cell_text(cell, value: str, *, bold: bool = False, color: str = "24364B", size: float = 9.0, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    shade(cell, "EAF2F8" if bold else "FFFFFF")
    borders(cell)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("V5 个人项目报告 | ")
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(8.5)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def set_run(run, *, size: float, color: str = "24364B", bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_para(doc: Document, text: str, *, size: float = 10.5, color: str = "24364B", bold: bool = False, align=None, after: float = 7, before: float = 0) -> None:
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.line_spacing = 1.45
    run = paragraph.add_run(text)
    set_run(run, size=size, color=color, bold=bold)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run(run, size={1: 17, 2: 13, 3: 11}[level], color="17324D", bold=True)


def add_callout(doc: Document, title: str, body: str, fill: str = "EFF6FA") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    shade(cell, fill)
    borders(cell, "B9D2E6")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    label = paragraph.add_run(title + "  ")
    set_run(label, size=10.5, color="17324D", bold=True)
    content = paragraph.add_run(body)
    set_run(content, size=10.5, color="24364B")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullets(doc: Document, rows: list[str]) -> None:
    for item in rows:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.35
        run = paragraph.add_run(item)
        set_run(run, size=10.2)


def add_figure(doc: Document, path: Path, caption: str, width_cm: float = 16.4) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.line_spacing = 1.1
    run = cap.add_run(caption)
    set_run(run, size=8.6, color="5E6F82")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        set_cell_text(cell, value, bold=True, color="17324D", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = Cm(widths[index])
            align = WD_ALIGN_PARAGRAPH.CENTER if index else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[index], value, size=8.5, align=align)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for level in (1, 2, 3):
        style = styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(23, 50, 77)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(16 if level == 1 else 11)
        style.paragraph_format.space_after = Pt(7)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("V5 | 金融科技与统计个人项目")
    set_run(header_run, size=8.5, color="5E6F82", bold=True)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run("V5")
    set_run(run, size=34, color="17324D", bold=True)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run("从银行多因子价值策略到\n个人化红利增强 ETF 的 Agent 研究流程")
    set_run(run, size=22, color="17324D", bold=True)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(20)
    run = paragraph.add_run("金融科技与统计个人项目")
    set_run(run, size=13, color="3F8F98", bold=True)
    add_callout(
        doc,
        "报告定位",
        "本报告展示一个以银行板块为起点、以理论解释和统计验证共同约束的量化研究流程。它不是实盘策略说明书、募资材料、投资建议或收益承诺。",
        "F3F8FB",
    )
    add_para(doc, "研究者：个人独立研究项目", size=10.5, color="5E6F82", align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    add_para(doc, "数据边界：不使用 2026-05-31 之后的市场数据", size=10.5, color="5E6F82", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.add_page_break()


def markdown() -> str:
    return f"""# V5：从银行多因子价值策略到个人化红利增强 ETF 的 Agent 研究流程

**金融科技与统计个人项目**

## 摘要

V5 是一个从 V1-V4 银行多因子价值策略出发的可审计量化研究项目。它的核心不是把历史收益包装成结论，而是把金融理论、点时可见数据、统计验证、工程复现和项目治理拆成可追溯的协作流程。研究将银行方法迁移到高速基础设施、港口铁路基础设施和电力公用事业，与银行共同构成四个核心 sleeve；价值低波负责选股基础，固定的同 sleeve 动量治理只作为辅助。均值回归、日内技术信号和若干行业扩展均被保留为研究线或被拒绝，未被包装成可比较的收益模型。

唯一允许进入正式表现正文的比较为 `{BASELINE}` 与 `{CANDIDATE}`：共同样本为 {FORMAL_START} 至 {FORMAL_END} 的 {OBS} 个交易日。候选总收益为 120.69%，基线为 109.25%；但这只是历史、非独立观察，候选状态仍为 `primary_forward_paper_candidate_not_accepted`。

## 研究问题

如何把一个在银行板块已形成基础的多因子价值框架，转化为可迁移、可审计、能暴露失败原因的 agent 协作研究流程，并在多个近似红利板块中构建个人化的红利增强 ETF 原型？

## 方法和结果

本报告使用固定时间边界：2013-01-01 至 2021-04-30 为研究与验证证据区间，2021-05-01 至 2026-05-31 为冻结的正式历史回测窗口。V5 四 sleeve 的完整 pre-2021 exact PIT 重建仍不完整，因此不把早期材料写成已完成的独立验证。正式回测样本首个有效信号日为 2021-05-06。

正式基线含银行、高速基础设施、港口铁路基础设施和电力公用事业四个核心 sleeve。主候选只在同 sleeve 内以固定 70/30 结构治理 12-1 动量，不改变价值低波选股池。它在正式共同样本中表现更高，但严格现金 NAV、QMT 历史合同、ETF total-return 合同和独立前瞻纸面周期均未形成，因此不具备 accepted 或实盘资格。

## 限制

- `strict_cash_nav_unavailable`
- `qmt_contract_incomplete`
- ETF total-return 合同缺失
- `validation_not_independent`
- 独立前瞻纸面周期尚未形成

这些限制不阻塞诚实的历史研究报告，但阻塞 accepted、live、部署批准和强收益结论。
"""


def build_pdf() -> Path:
    pdfmetrics.registerFont(TTFont("MSYH", r"C:\Windows\Fonts\msyh.ttc"))
    pdfmetrics.registerFont(TTFont("MSYHB", r"C:\Windows\Fonts\msyhbd.ttc"))
    pdf = OUT / "v5x_fintech_statistics_personal_project_report.pdf"
    navy, ink, muted, teal = "#17324D", "#24364B", "#5E6F82", "#3F8F98"
    styles = getSampleStyleSheet()
    body = ParagraphStyle("BodyCN", parent=styles["BodyText"], fontName="MSYH", fontSize=9.5, leading=15, textColor=colors.HexColor(ink), spaceAfter=8, wordWrap="CJK")
    h1 = ParagraphStyle("H1CN", parent=styles["Heading1"], fontName="MSYHB", fontSize=17, leading=23, textColor=colors.HexColor(navy), spaceBefore=14, spaceAfter=8, keepWithNext=True, wordWrap="CJK")
    h2 = ParagraphStyle("H2CN", parent=styles["Heading2"], fontName="MSYHB", fontSize=12.5, leading=18, textColor=colors.HexColor(navy), spaceBefore=10, spaceAfter=6, keepWithNext=True, wordWrap="CJK")
    small = ParagraphStyle("SmallCN", parent=body, fontSize=8.2, leading=11.5, textColor=colors.HexColor(muted), alignment=TA_CENTER, spaceAfter=8, wordWrap="CJK")
    callout = ParagraphStyle("CalloutCN", parent=body, fontSize=9.3, leading=14.2, textColor=colors.HexColor(ink), spaceAfter=0, wordWrap="CJK")
    cover_title = ParagraphStyle("CoverTitle", parent=h1, fontName="MSYHB", fontSize=23, leading=32, textColor=colors.HexColor(navy), alignment=TA_CENTER, spaceAfter=13, wordWrap="CJK")
    cover_subtitle = ParagraphStyle("CoverSub", parent=body, fontName="MSYHB", fontSize=12.5, leading=18, textColor=colors.HexColor(teal), alignment=TA_CENTER, spaceAfter=20, wordWrap="CJK")

    def para(value: str, style=body):
        return Paragraph(value, style)

    def figure(path: Path, caption: str):
        image = PdfImage(str(path))
        max_width = 16.0 * cm
        ratio = image.imageHeight / image.imageWidth
        image.drawWidth = max_width
        image.drawHeight = max_width * ratio
        return KeepTogether([image, Spacer(1, 3), para(caption, small), Spacer(1, 5)])

    def callout_box(title: str, text: str, color="#F3F8FB"):
        table = Table([[para(f"<b>{title}</b>  {text}", callout)]], colWidths=[16.3 * cm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(color)),
            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#B9D2E6")),
            ("LEFTPADDING", (0, 0), (-1, -1), 11),
            ("RIGHTPADDING", (0, 0), (-1, -1), 11),
            ("TOPPADDING", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ]))
        return [table, Spacer(1, 10)]

    def data_table(headers: list[str], rows: list[list[str]], col_widths: list[float]):
        wrapped = [[para(value, ParagraphStyle("TableHead", parent=small, fontName="MSYHB", fontSize=8.0, leading=10, textColor=colors.HexColor(navy), alignment=TA_CENTER, spaceAfter=0)) for value in headers]]
        for row in rows:
            wrapped.append([para(value, ParagraphStyle("TableBody", parent=body, fontSize=8.0, leading=10.5, alignment=TA_CENTER if index else 0, spaceAfter=0, wordWrap="CJK")) for index, value in enumerate(row)])
        table = Table(wrapped, colWidths=[width * cm for width in col_widths], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF2F8")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D7E4EF")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        return [table, Spacer(1, 9)]

    def on_page(canvas, document):
        canvas.saveState()
        canvas.setFont("MSYH", 8)
        canvas.setFillColor(colors.HexColor(muted))
        canvas.drawString(1.8 * cm, 28.5 * cm, "V5 | 金融科技与统计个人项目")
        canvas.drawRightString(19.2 * cm, 1.15 * cm, f"V5 个人项目报告 | {document.page}")
        canvas.restoreState()

    story = [Spacer(1, 4.5 * cm), para("V5", ParagraphStyle("V5", parent=cover_title, fontSize=34, leading=42)), para("从银行多因子价值策略到<br/>个人化红利增强 ETF 的 Agent 研究流程", cover_title), para("金融科技与统计个人项目", cover_subtitle)]
    story.extend(callout_box("报告定位", "本报告展示一个以银行板块为起点、以理论解释和统计验证共同约束的量化研究流程。它不是实盘策略说明书、募资材料、投资建议或收益承诺文件。"))
    story += [para("研究者：个人独立研究项目", small), para("数据边界：不使用 2026-05-31 之后的市场数据", small), PageBreak()]

    story += [para("执行摘要", h1), para("V5 以 V1-V4 的银行多因子价值策略为研究起点，尝试解决一个比“找出更高历史收益”更基础的问题：如何把金融理论、点时可见数据、统计检验、工程实现和治理判断组织为一套能被他人复核的量化研究流程。"), para("研究最终形成四个核心 sleeve：银行、高速基础设施、港口铁路基础设施和电力公用事业。价值低波是共同的选股骨架；同 sleeve 的固定动量治理是唯一进入正式 pairwise 比较的辅助机制。均值回归、分钟级执行与技术分析用于检验边界和执行假设，并未因局部结果被升级为可投资策略。")]
    story.extend(callout_box("正式结果的正确读法", "在 2021-05-06 至 2026-05-29 的 1,228 个共同交易日中，主候选的历史总收益为 120.69%，repaired baseline 为 109.25%。这是同父策略、同回报合同下的非独立历史观察，不是独立验证、收益预测或实盘批准。", "#FFF8ED"))

    story += [para("一、项目问题与个人贡献", h1), para("本项目把银行板块中已经经过多轮统计与经济解释检验的价值策略，作为一个可复制的研究起点，而非把它当作可以直接套用的答案。研究目标是建立一个面向近似红利板块的 agent 开发工作流：先提出有金融解释的假设，再由量化程序进行反证、归因和分层，最后由工程与项目治理判断它是否值得被保留。")]
    for value in ["研究设计：定义时间序列隔离、PIT 数据可见性和禁止事后调参的边界。", "理论到检验：将银行策略的价值、现金流质量、低波动与分红逻辑拆解为可检验假设。", "多 agent 协作：研究员、量化、工程师和项目经理的职责互相制约；顾问对话保持问题定义与解释链条。", "失败可追溯：把未通过的行业、均值回归、技术分析和执行机制保留为被拒绝或观察的研究证据。"]:
        story.append(para("• " + value))
    story += [figure(FIG_V / "v5_core_purpose_overview.png", "图 1. V5 的核心目的：从银行策略基础到多 sleeve、可迁移的研究流程。"), figure(FIG_V / "v5_agent_workflow_overview.png", "图 2. 研究员、量化、工程、项目经理与顾问之间的协作和监督关系。")]

    story += [para("二、研究边界：时间隔离与点时可见原则", h1), para("时间序列研究最容易出现的错误，是把未来信息或反复试验后的规则混入历史评价。V5 将 2013-01-01 至 2021-04-30 定义为研究与验证证据区间，将 2021-05-01 至 2026-05-31 固定为正式历史回测窗口。对银行主线，早期材料支持滚动式研究；但四 sleeve 的完整 pre-2021 exact PIT 重建尚未完成，因此不能被写作跨周期的独立验证。")]
    story.extend(callout_box("PIT 约束", "财报、行业纳入、分红与公司行为必须以当时可见的日期进入研究面板。无法以原始证据闭合的字段被保留为数据门或限制，而不以未来年报、代理目标或人工估计回填。"))
    story += [figure(FIG_V / "v5_time_series_sample_isolation_design.png", "图 3. 时间序列隔离设计：早期研究证据、冻结回测窗口与防止样本外污染的边界。"), para("三、理论、统计与工程的共同约束", h1), para("单纯的数学优化可以给出一个历史上看似更优的组合，却未必能解释为何它应当存在。V5 的研究顺序因此反过来：先以价值投资、现金流质量、行业经营模式和市场情绪为理论来源；再让量化 agent 通过截面检验、样本切分、失败归因和回测合同去反驳它；最后让工程 agent 检验数据、执行与现金路径是否真的可复现。")]
    story += [figure(FIG_V / "v5_finance_theory_statistics_agent_loop.png", "图 4. 金融理论提出假设，统计检验尝试证伪，工程与项目治理决定是否可保留。"), figure(FIG_V / "v5_ideal_vs_reality_optimization_path.png", "图 5. 理想优化与实际约束：理论解释、统计证据与工程实现必须同时面对数据边界。")]

    story += [PageBreak(), para("四、从银行方法到四个核心 sleeve", h1), para("正式 baseline 为 startup preload 修复后的 V57f：银行、高速基础设施、港口铁路基础设施和电力公用事业共同构成四个核心 sleeve。它们不是按事后收益排名入选，而是基于经营模式、现金流和分红逻辑的可解释性、PIT 资料可用性及逐步通过的数据门。"), para("行业扩展也说明了迁移边界。电信、燃气水务和保险等仅保留为观察或专项研究；煤炭等周期行业则因周期状态与 PIT 数据门不足而阻塞。研究的结论不是“所有近似行业都应被纳入”，而是流程可以清楚地区分哪些路径有证据、哪些仍需研究、哪些不应升级。")]
    story += [figure(FIG_W / "v5w_sector_research_decision_funnel.png", "图 6. 八个行业路径的研究漏斗：筛选、观察、数据门和未升级原因被明确记录。"), figure(FIG_V / "v5_core_sleeve_reference_comparison_2013_2026.png", "图 7. 四个核心 sleeve 的市场背景轨迹，仅用于说明研究背景，不构成正式绩效比较。")]

    story += [PageBreak(), para("五、正式基线与主候选：唯一允许的表现比较", h1), para(f"唯一正式基线为 {BASELINE}，唯一主候选为 {CANDIDATE}。二者的比较限定为 {FORMAL_START} 至 {FORMAL_END}、{OBS} 个共同交易日、同父策略和同一日度回报合同。候选并不改变价值低波选股池，只在同 sleeve 内以预先固定的 70/30 方式处理 12-1 动量治理。")]
    story += data_table(["模型", "总收益", "年化收益", "最大回撤", "夏普", "相对基线"], [["repaired baseline", "109.25%", "16.36%", "11.93%", "1.048", "-"], ["主候选：内部子 sleeve 动量 70/30", "120.69%", "17.64%", "11.88%", "1.115", "+11.43 个百分点"]], [3.9, 2.0, 2.0, 1.8, 1.45, 2.25])
    story += [para("候选在该历史共同样本中表现更高，最大回撤差为 -0.05 个百分点，属于轻微改善。该结论的含义被严格限定为历史样本中的 pairwise observation；它不能被解读为更稳健、已验证或可部署的承诺。", ParagraphStyle("Note", parent=body, fontSize=8.8, leading=13, textColor=colors.HexColor(muted))), figure(FIG_V / "v5v_formal_pairwise_comparison.png", "图 8. repaired baseline 与主候选的正式同合同比较。"), figure(FIG_W / "v5w_formal_pairwise_annual_return_and_underwater.png", "图 9. 年度描述性拆解与共同窗口回撤路径；2021 和 2026 为非完整自然年。")]

    story += [PageBreak(), para("六、未升级研究线：结果之外的证据", h1), para("V5 没有把所有尝试都写成成功。均值回归在红利型板块中未形成足够稳定的增量证据；分钟级数据帮助诊断买入时点和执行路径，却没有自动转化为更高频交易规则；技术分析的卖出假设在工程验证后未形成可升级的统一规则。它们被保留为诊断、执行或观察材料，不能与正式基线并列排序。")]
    for value in ["均值回归：从“低估优质股”到分钟级急跌修复均进行了假设检验，但未获得足以进入主线的跨期证据。", "分钟执行：1 分钟和 5 分钟价格、成交量用于提高压力和 VWAP 等诊断精度，而不是提高策略交易频率。", "技术分析：卖出时点、日线价格与成交量等规则经过分阶段检验；因跨期或执行证据不足，未改写价值主线。", "行业扩展：研究线进入观察或数据门，不因为局部收益而自动并入核心 basket。"]:
        story.append(para("• " + value))
    story += [para("七、可复现性、现实摩擦与治理状态", h1), para("V5 把可复现性理解为完整的研究合同，而不只是能再次跑出一条净值曲线。每个正式材料应能追溯到代码、输入、配置、统计定义、生成时间和 SHA256 记录。平台回放、订单治理和桥接工程同样被保存，但它们的证据层级与历史策略表现不同。")]
    story.extend(callout_box("当前治理状态", "internal_subsleeve_mom12_70_30 的唯一状态是 primary_forward_paper_candidate_not_accepted。任何 accepted、live、deployment approved 或独立验证通过的表述均不成立。", "#FFF8ED"))
    story += data_table(["限制 / 证据缺口", "对报告的影响", "对升级的影响"], [["strict_cash_nav_unavailable", "不否定诚实的历史研究", "阻塞严格现金 NAV 与强执行结论"], ["qmt_contract_incomplete", "QMT 仅能作为工程/证据附录", "阻塞精确成交归因和实盘升级"], ["ETF total-return 合同缺失", "ETF 只作经济暴露背景", "禁止正式 ETF Alpha、Beta、IR 与超额收益"], ["validation_not_independent", "结果只能称历史观察", "阻塞独立验证结论"], ["独立前瞻纸面周期尚未形成", "历史报告可归档", "阻塞候选 accepted 与 live 语言"]], [5.1, 5.6, 5.5])

    story += [PageBreak(), para("八、结论与下一门槛", h1), para("V5 的主要产出不是一个可以被简单宣传为“已验证策略”的模型，而是一套可审计的研究能力：从银行价值策略的经济解释出发，使用点时可见数据和统计检验探索近似行业，记录工程可行性与失败原因，并通过项目治理阻止未成熟的研究线越过证据边界。"), para("在正式历史窗口内，主候选相对于 repaired baseline 出现了有限、可量化的正向观察；同时，项目完整保存了它尚不能升级的理由。下一步不是继续在旧样本上优化，而是在获得授权后，从首笔纸面订单起形成完整的 target -> order -> fill -> position -> cash 前瞻闭环，再按预先固定的门槛进行复核。")]
    story.extend(callout_box("个人项目的定位", "本项目展示的是金融理论、统计方法、数据工程和 agent 协作如何共同提高研究过程的透明度与可复现性；不构成投资建议、未来收益预测或实盘系统推荐。"))
    story += [para("附录：事实冻结与阅读说明", h1)]
    for value in ["唯一正式基线：v57f_startup_preload_repaired_baseline。", "正式窗口：2021-05-01 至 2026-05-31；首个有效信号日：2021-05-06。", "唯一正式比较样本：2021-05-06 至 2026-05-29，1,228 个共同交易日。", "ETF 图仅用于 economic exposure context，不使用 adjusted price return 计算正式 Alpha、Beta、IR 或超额收益。", "V5c 诊断、V5d/V5h/V5i/V5j 执行与技术研究、V5e 现金研究、QMT 回放和订单治理均不进入策略收益排序。", "材料追溯：V5t 历史研究与治理报告 v1.0、V5s 数字复核、V5w 报告准备包及其 SHA256 manifest。"]:
        story.append(para("• " + value))

    document = SimpleDocTemplate(str(pdf), pagesize=A4, rightMargin=1.8 * cm, leftMargin=1.8 * cm, topMargin=2.0 * cm, bottomMargin=1.8 * cm, title="V5 金融科技与统计研究项目")
    document.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return pdf


def build_report() -> tuple[Path, Path]:
    doc = Document()
    configure(doc)
    add_cover(doc)

    add_heading(doc, "执行摘要")
    add_para(doc, "V5 以 V1-V4 的银行多因子价值策略为研究起点，尝试解决一个比“找出更高历史收益”更基础的问题：如何把金融理论、点时可见数据、统计检验、工程实现和治理判断组织为一套能被他人复核的量化研究流程。")
    add_para(doc, "研究最终形成四个核心 sleeve：银行、高速基础设施、港口铁路基础设施和电力公用事业。价值低波是共同的选股骨架；同 sleeve 的固定动量治理是唯一进入正式 pairwise 比较的辅助机制。均值回归、分钟级执行与技术分析用于检验边界和执行假设，并未因局部结果被升级为可投资策略。")
    add_callout(doc, "正式结果的正确读法", "在 2021-05-06 至 2026-05-29 的 1,228 个共同交易日中，主候选的历史总收益为 120.69%，repaired baseline 为 109.25%。这是同父策略、同回报合同下的非独立历史观察，不是独立验证、收益预测或实盘批准。", "FFF8ED")

    add_heading(doc, "一、项目问题与个人贡献")
    add_para(doc, "本项目把银行板块中已经经过多轮统计与经济解释检验的价值策略，作为一个可复制的研究起点，而非把它当作可以直接套用的答案。研究目标是建立一个面向近似红利板块的 agent 开发工作流：先提出有金融解释的假设，再由量化程序进行反证、归因和分层，最后由工程与项目治理判断它是否值得被保留。")
    add_bullets(doc, [
        "研究设计：定义时间序列隔离、PIT 数据可见性和禁止事后调参的边界。",
        "理论到检验：将银行策略的价值、现金流质量、低波动与分红逻辑拆解为可检验假设。",
        "多 agent 协作：研究员、量化、工程师和项目经理的职责互相制约；顾问对话保持问题定义与解释链条。",
        "失败可追溯：把未通过的行业、均值回归、技术分析和执行机制保留为被拒绝或观察的研究证据。",
    ])
    add_figure(doc, FIG_V / "v5_core_purpose_overview.png", "图 1. V5 的核心目的：从银行策略基础到多 sleeve、可迁移的研究流程。")
    add_figure(doc, FIG_V / "v5_agent_workflow_overview.png", "图 2. 研究员、量化、工程、项目经理与顾问之间的协作和监督关系。")

    add_heading(doc, "二、研究边界：时间隔离与点时可见原则")
    add_para(doc, "时间序列研究最容易出现的错误，是把未来信息或反复试验后的规则混入历史评价。V5 将 2013-01-01 至 2021-04-30 定义为研究与验证证据区间，将 2021-05-01 至 2026-05-31 固定为正式历史回测窗口。对银行主线，早期材料支持滚动式研究；但四 sleeve 的完整 pre-2021 exact PIT 重建尚未完成，因此不能被写作跨周期的独立验证。")
    add_callout(doc, "PIT 约束", "财报、行业纳入、分红与公司行为必须以当时可见的日期进入研究面板。无法以原始证据闭合的字段被保留为数据门或限制，而不以未来年报、代理目标或人工估计回填。")
    add_figure(doc, FIG_V / "v5_time_series_sample_isolation_design.png", "图 3. 时间序列隔离设计：早期研究证据、冻结回测窗口与防止样本外污染的边界。")

    add_heading(doc, "三、理论、统计与工程的共同约束")
    add_para(doc, "单纯的数学优化可以给出一个历史上看似更优的组合，却未必能解释为何它应当存在。V5 的研究顺序因此反过来：先以价值投资、现金流质量、行业经营模式和市场情绪为理论来源；再让量化 agent 通过截面检验、样本切分、失败归因和回测合同去反驳它；最后让工程 agent 检验数据、执行与现金路径是否真的可复现。")
    add_para(doc, "项目经理不把“收益更高”视为唯一判定，而是区分三种失败：金融解释不足、统计证据不足、或工程成本和数据边界不允许可靠实现。这一过程使未通过的路径同样具有研究价值。")
    add_figure(doc, FIG_V / "v5_finance_theory_statistics_agent_loop.png", "图 4. 金融理论提出假设，统计检验尝试证伪，工程与项目治理决定是否可保留。")
    add_figure(doc, FIG_V / "v5_ideal_vs_reality_optimization_path.png", "图 5. 理想优化与实际约束：理论解释、统计证据与工程实现必须同时面对数据边界。")

    add_heading(doc, "四、从银行方法到四个核心 sleeve")
    add_para(doc, "正式 baseline 为 startup preload 修复后的 V57f：银行、高速基础设施、港口铁路基础设施和电力公用事业共同构成四个核心 sleeve。它们不是按事后收益排名入选，而是基于经营模式、现金流和分红逻辑的可解释性、PIT 资料可用性及逐步通过的数据门。")
    add_para(doc, "行业扩展也说明了迁移边界。电信、燃气水务和保险等仅保留为观察或专项研究；煤炭等周期行业则因周期状态与 PIT 数据门不足而阻塞。研究的结论不是“所有近似行业都应被纳入”，而是流程可以清楚地区分哪些路径有证据、哪些仍需研究、哪些不应升级。")
    add_figure(doc, FIG_W / "v5w_sector_research_decision_funnel.png", "图 6. 八个行业路径的研究漏斗：筛选、观察、数据门和未升级原因被明确记录。")
    add_figure(doc, FIG_V / "v5_core_sleeve_reference_comparison_2013_2026.png", "图 7. 四个核心 sleeve 的市场背景轨迹，仅用于说明研究背景，不构成正式绩效比较。")

    add_heading(doc, "五、正式基线与主候选：唯一允许的表现比较")
    add_para(doc, f"唯一正式基线为 `{BASELINE}`，唯一主候选为 `{CANDIDATE}`。二者的比较限定为 {FORMAL_START} 至 {FORMAL_END}、{OBS} 个共同交易日、同父策略和同一日度回报合同。候选并不改变价值低波选股池，只在同 sleeve 内以预先固定的 70/30 方式处理 12-1 动量治理。")
    add_table(doc,
        ["模型", "总收益", "年化收益", "最大回撤", "夏普", "相对基线"],
        [
            ["repaired baseline", "109.25%", "16.36%", "11.93%", "1.048", "-"],
            ["主候选：内部子 sleeve 动量 70/30", "120.69%", "17.64%", "11.88%", "1.115", "+11.43 个百分点"],
        ],
        [4.0, 2.1, 2.1, 2.0, 1.55, 2.2],
    )
    add_para(doc, "候选在该历史共同样本中表现更高，最大回撤差为 -0.05 个百分点，属于轻微改善。该结论的含义被严格限定为历史样本中的 pairwise observation；它不能被解读为更稳健、已验证或可部署的承诺。", size=9.8, color="5E6F82")
    add_figure(doc, FIG_V / "v5v_formal_pairwise_comparison.png", "图 8. repaired baseline 与主候选的正式同合同比较。")
    add_figure(doc, FIG_W / "v5w_formal_pairwise_annual_return_and_underwater.png", "图 9. 年度描述性拆解与共同窗口回撤路径；2021 和 2026 为非完整自然年。")

    add_heading(doc, "六、未升级研究线：结果之外的证据")
    add_para(doc, "V5 没有把所有尝试都写成成功。均值回归在红利型板块中未形成足够稳定的增量证据；分钟级数据帮助诊断买入时点和执行路径，却没有自动转化为更高频交易规则；技术分析的卖出假设在工程验证后未形成可升级的统一规则。它们被保留为诊断、执行或观察材料，不能与正式基线并列排序。")
    add_bullets(doc, [
        "均值回归：从“低估优质股”到分钟级急跌修复均进行了假设检验，但未获得足以进入主线的跨期证据。",
        "分钟执行：1 分钟和 5 分钟价格、成交量用于提高压力和 VWAP 等诊断精度，而不是提高策略交易频率。",
        "技术分析：卖出时点、日线价格与成交量等规则经过分阶段检验；因跨期或执行证据不足，未改写价值主线。",
        "行业扩展：研究线进入观察或数据门，不因为局部收益而自动并入核心 basket。",
    ])

    add_heading(doc, "七、可复现性、现实摩擦与治理状态")
    add_para(doc, "V5 把可复现性理解为完整的研究合同，而不只是能再次跑出一条净值曲线。每个正式材料应能追溯到代码、输入、配置、统计定义、生成时间和 SHA256 记录。平台回放、订单治理和桥接工程同样被保存，但它们的证据层级与历史策略表现不同。")
    add_callout(doc, "当前治理状态", "`internal_subsleeve_mom12_70_30` 的唯一状态是 `primary_forward_paper_candidate_not_accepted`。任何 accepted、live、deployment approved 或独立验证通过的表述均不成立。", "FFF8ED")
    add_table(doc,
        ["限制 / 证据缺口", "对报告的影响", "对升级的影响"],
        [
            ["strict_cash_nav_unavailable", "不否定诚实的历史研究", "阻塞严格现金 NAV 与强执行结论"],
            ["qmt_contract_incomplete", "QMT 仅能作为工程/证据附录", "阻塞精确成交归因和实盘升级"],
            ["ETF total-return 合同缺失", "ETF 只作经济暴露背景", "禁止正式 ETF Alpha、Beta、IR 与超额收益"],
            ["validation_not_independent", "结果只能称历史观察", "阻塞独立验证结论"],
            ["独立前瞻纸面周期尚未形成", "历史报告可归档", "阻塞候选 accepted 与 live 语言"],
        ],
        [4.0, 4.1, 4.0],
    )

    add_heading(doc, "八、结论与下一门槛")
    add_para(doc, "V5 的主要产出不是一个可以被简单宣传为“已验证策略”的模型，而是一套可审计的研究能力：从银行价值策略的经济解释出发，使用点时可见数据和统计检验探索近似行业，记录工程可行性与失败原因，并通过项目治理阻止未成熟的研究线越过证据边界。")
    add_para(doc, "在正式历史窗口内，主候选相对于 repaired baseline 出现了有限、可量化的正向观察；同时，项目完整保存了它尚不能升级的理由。下一步不是继续在旧样本上优化，而是在获得授权后，从首笔纸面订单起形成完整的 `target -> order -> fill -> position -> cash` 前瞻闭环，再按预先固定的门槛进行复核。")
    add_callout(doc, "个人项目的定位", "本项目展示的是金融理论、统计方法、数据工程和 agent 协作如何共同提高研究过程的透明度与可复现性；不构成投资建议、未来收益预测或实盘系统推荐。", "F3F8FB")

    add_heading(doc, "附录：事实冻结与阅读说明")
    add_bullets(doc, [
        "唯一正式基线：v57f_startup_preload_repaired_baseline。",
        "正式窗口：2021-05-01 至 2026-05-31；首个有效信号日：2021-05-06。",
        "唯一正式比较样本：2021-05-06 至 2026-05-29，1,228 个共同交易日。",
        "ETF 图仅用于 economic exposure context，不使用 adjusted price return 计算正式 Alpha、Beta、IR 或超额收益。",
        "V5c 诊断、V5d/V5h/V5i/V5j 执行与技术研究、V5e 现金研究、QMT 回放和订单治理均不进入策略收益排序。",
        "材料追溯：V5t 历史研究与治理报告 v1.0、V5s 数字复核、V5w 报告准备包及其 SHA256 manifest。",
    ])
    docx = OUT / "v5x_fintech_statistics_personal_project_report.docx"
    doc.save(docx)
    md = OUT / "v5x_fintech_statistics_personal_project_report.md"
    md.write_text(markdown(), encoding="utf-8")
    return docx, md


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    docx, md = build_report()
    pdf = build_pdf()
    required = [
        ROOT / "v5t_report_publication_finalization" / "current" / "v5t_publication_summary.json",
        ROOT / "v5s_report_draft_review_and_release_prep" / "current" / "v5s_numeric_reconciliation.csv",
        ROOT / "v5v_application_final_report_readiness" / "current" / "v5v_application_narrative_freeze.md",
        ROOT / "v5v_application_final_report_readiness" / "current" / "v5v_figure_register.csv",
        ROOT / "v5w_final_report_preparation" / "current" / "v5w_final_report_preparation_summary.json",
        ROOT / "v5w_final_report_preparation" / "current" / "v5w_source_contract_conflict_log.csv",
        ROOT / "docs" / "governance" / "status_registry.json",
    ]
    figures = [
        FIG_V / "v5_core_purpose_overview.png",
        FIG_V / "v5_agent_workflow_overview.png",
        FIG_V / "v5_time_series_sample_isolation_design.png",
        FIG_V / "v5_finance_theory_statistics_agent_loop.png",
        FIG_V / "v5_ideal_vs_reality_optimization_path.png",
        FIG_W / "v5w_sector_research_decision_funnel.png",
        FIG_V / "v5_core_sleeve_reference_comparison_2013_2026.png",
        FIG_V / "v5v_formal_pairwise_comparison.png",
        FIG_W / "v5w_formal_pairwise_annual_return_and_underwater.png",
    ]
    missing = [str(path) for path in required + figures if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing frozen report inputs: " + "; ".join(missing))
    git_head = subprocess.run(["git", "rev-parse", "HEAD"], cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
    rows = []
    for path in required + figures:
        rows.append({"path": str(path.relative_to(ROOT)), "sha256": sha256(path), "role": "frozen_report_input"})
    rows += [
        {"path": str(docx.relative_to(ROOT)), "sha256": sha256(docx), "role": "generated_docx"},
        {"path": str(md.relative_to(ROOT)), "sha256": sha256(md), "role": "generated_markdown"},
        {"path": str(pdf.relative_to(ROOT)), "sha256": sha256(pdf), "role": "generated_pdf"},
    ]
    manifest = OUT / "v5x_personal_project_report_input_manifest.csv"
    with manifest.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["path", "sha256", "role"])
        writer.writeheader()
        writer.writerows(rows)
    usage = OUT / "v5x_personal_project_figure_usage_register.csv"
    with usage.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["figure", "report_section", "permitted_use", "boundary"])
        writer.writeheader()
        for path, section, permitted, boundary in [
            ("v5_core_purpose_overview.png", "项目问题与个人贡献", "project overview", "no performance claim"),
            ("v5_agent_workflow_overview.png", "项目问题与个人贡献", "method and governance", "no performance claim"),
            ("v5_time_series_sample_isolation_design.png", "研究边界", "study design", "no performance claim"),
            ("v5_finance_theory_statistics_agent_loop.png", "理论、统计与工程", "research method", "no performance claim"),
            ("v5_ideal_vs_reality_optimization_path.png", "理论、统计与工程", "reflection and limitations", "no performance claim"),
            ("v5w_sector_research_decision_funnel.png", "四个核心 sleeve", "research route", "not a return ranking"),
            ("v5_core_sleeve_reference_comparison_2013_2026.png", "四个核心 sleeve", "economic context", "not formal performance"),
            ("v5v_formal_pairwise_comparison.png", "正式基线与主候选", "formal pairwise", "validation_not_independent"),
            ("v5w_formal_pairwise_annual_return_and_underwater.png", "正式基线与主候选", "formal pairwise", "validation_not_independent"),
        ]:
            writer.writerow({"figure": path, "report_section": section, "permitted_use": permitted, "boundary": boundary})
    fact_freeze = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "report_kind": "personal_fintech_statistics_research_project",
        "baseline": BASELINE,
        "primary_candidate": CANDIDATE,
        "candidate_status": "primary_forward_paper_candidate_not_accepted",
        "formal_window": "2021-05-01_to_2026-05-31",
        "formal_common_sample": {"start": FORMAL_START, "end": FORMAL_END, "observations": OBS},
        "formal_performance_scope": "baseline_vs_primary_candidate_only",
        "required_disclosures": [
            "strict_cash_nav_unavailable",
            "qmt_contract_incomplete",
            "ETF_total_return_contract_missing",
            "validation_not_independent",
            "independent_forward_paper_cycle_not_formed",
        ],
        "git_head": git_head,
        "market_data_scope_end": "2026-05-31",
    }
    (OUT / "v5x_personal_project_fact_freeze.json").write_text(json.dumps(fact_freeze, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"docx": str(docx), "markdown": str(md), "pdf": str(pdf), "manifest": str(manifest)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
